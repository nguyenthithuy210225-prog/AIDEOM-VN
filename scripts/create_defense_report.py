# -*- coding: utf-8 -*-
"""Create the AIDEOM-VN defense report DOCX."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "BAO_CAO_BAO_VE_AIDEOM_VN.docx"
FIG = ROOT / "outputs" / "figures"
FONT = "Arial"
DARK = RGBColor(17, 24, 39)
TEAL = RGBColor(0, 128, 112)
GRAY = RGBColor(75, 85, 99)
BORDER = "CBD5E1"


def _rfonts_for_run(run, font=FONT):
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(key), font)


def _rfonts_for_style(style, font=FONT):
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(key), font)


def style_font(style, size=None, bold=None, color=None):
    style.font.name = FONT
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    _rfonts_for_style(style)


def add_run(paragraph, text, size=None, bold=None, color=None, italic=None):
    run = paragraph.add_run(text)
    run.font.name = FONT
    _rfonts_for_run(run)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell(cell, text, bold=False, color=DARK, fill=None, align=None):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        shade_cell(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align or WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    add_run(paragraph, str(text), size=9.2, bold=bold, color=color)


def border_table(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    paragraph.style = f"Heading {level}"
    paragraph.paragraph_format.space_before = Pt(8 if level == 1 else 4)
    paragraph.paragraph_format.space_after = Pt(6)
    add_run(
        paragraph,
        text,
        size=18 if level == 1 else 13,
        bold=True,
        color=TEAL if level == 1 else DARK,
    )
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.12
    add_run(paragraph, text, size=10.4, color=DARK)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    add_run(paragraph, text, size=10.1, color=DARK)
    return paragraph


def add_callout(doc, title, bullets):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    border_table(table)
    cell = table.cell(0, 0)
    shade_cell(cell, "ECFDF5")
    cell.paragraphs[0].clear()
    add_run(cell.paragraphs[0], title, size=11, bold=True, color=TEAL)
    for bullet in bullets:
        paragraph = cell.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        add_run(paragraph, bullet, size=9.7, color=DARK)
    doc.add_paragraph()


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    border_table(table)
    for index, header in enumerate(headers):
        set_cell(
            table.rows[0].cells[index],
            header,
            bold=True,
            color=RGBColor(255, 255, 255),
            fill="0F766E",
        )
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell(cells[index], value)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    doc.add_paragraph()
    return table


def add_picture_if_exists(doc, filename, caption, width=6.2):
    path = FIG / filename
    if not path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(caption_paragraph, caption, size=8.8, italic=True, color=GRAY)


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    for style_name, size, bold, color in [
        ("Normal", 10.4, False, DARK),
        ("Heading 1", 18, True, TEAL),
        ("Heading 2", 13, True, DARK),
        ("List Bullet", 10.1, False, DARK),
        ("Title", 26, True, DARK),
    ]:
        style_font(doc.styles[style_name], size, bold, color)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(footer, "AIDEOM-VN Decision Lab | Báo cáo bảo vệ cuối kỳ", size=8.5, color=GRAY)
    return doc


def add_cover(doc):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(42)
    add_run(paragraph, "AIDEOM-VN DECISION LAB", size=15, bold=True, color=TEAL)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(
        paragraph,
        "WEB APP MÔ HÌNH RA QUYẾT ĐỊNH\nPHÁT TRIỂN KINH TẾ VIỆT NAM",
        size=24,
        bold=True,
        color=DARK,
    )

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(
        paragraph,
        "Báo cáo giới thiệu dự án, mô tả chức năng và hướng dẫn sử dụng",
        size=12,
        color=GRAY,
    )

    add_table(
        doc,
        ["Hạng mục", "Nội dung"],
        [
            ["Tên sản phẩm", "AIDEOM-VN Policy Cockpit"],
            ["Hình thức", "Web app Streamlit chạy local, có thể đóng gói nộp kèm dữ liệu và output"],
            ["Phạm vi", "12 bài/mô-đun quyết định theo đề cuối kỳ"],
            ["Dữ liệu", "CSV local 2020-2025, dữ liệu ngành, vùng, dự án và các bảng mô phỏng phụ trợ"],
            ["Đầu ra", "Bảng CSV, biểu đồ PNG, dashboard tương tác và tác nhân phân tích kết quả"],
        ],
    )
    add_callout(
        doc,
        "Thông điệp bảo vệ",
        [
            "Dự án tích hợp các bài toán định lượng thành một phòng thí nghiệm ra quyết định có tham số, biểu đồ, bảng kết quả và diễn giải chuyên gia.",
            "Toàn bộ pipeline có thể tái lập trên máy local: đọc dữ liệu, tính toán, sinh bảng, sinh hình và chạy giao diện.",
        ],
    )
    doc.add_page_break()


def add_toc(doc):
    add_heading(doc, "Mục lục", 1)
    for item in [
        "1. Tổng quan dự án",
        "2. Kiến trúc hệ thống và dữ liệu",
        "3. Mô tả chi tiết 12 chức năng",
        "4. Tác nhân phân tích kết quả",
        "5. Hướng dẫn cài đặt và sử dụng",
        "6. Kết quả đầu ra và kiểm thử",
        "7. Kết luận",
    ]:
        add_body(doc, item)
    doc.add_page_break()


def add_report_content(doc):
    add_heading(doc, "1. Tổng quan dự án", 1)
    add_body(
        doc,
        "AIDEOM-VN là web app mô phỏng và tối ưu hóa quyết định phát triển kinh tế Việt Nam trong bối cảnh chuyển đổi số và AI. Ứng dụng được xây dựng theo phương án web app tương tác: người dùng chọn từng bài ở menu, chỉnh tham số, xem bảng kết quả, biểu đồ, giải thích mô hình và tải output phục vụ báo cáo.",
    )
    add_body(
        doc,
        "Điểm mạnh của sản phẩm là đưa các bài toán định lượng vào cùng một dashboard thống nhất. Các mô hình từ Cobb-Douglas, quy hoạch tuyến tính, MCDM, MIP, TOPSIS, Pareto, tối ưu động, stochastic programming đến Q-learning đều có giao diện riêng và có diễn giải kết quả ở cuối mỗi bài.",
    )
    add_table(
        doc,
        ["Nhóm chức năng", "Mục tiêu"],
        [
            ["Mô hình dự báo", "Ước lượng TFP, dự báo GDP và phân rã đóng góp tăng trưởng"],
            ["Tối ưu ngân sách", "Phân bổ nguồn lực theo ràng buộc ngân sách, rủi ro, vùng và ngành"],
            ["Xếp hạng ưu tiên", "Chuẩn hóa tiêu chí, TOPSIS và phân tích độ nhạy"],
            ["Mô phỏng chính sách", "So sánh kịch bản dài hạn, tác động lao động và học tăng cường"],
            ["Báo cáo", "Sinh bảng CSV, biểu đồ PNG và diễn giải chuyên gia phục vụ bảo vệ"],
        ],
    )
    add_picture_if_exists(doc, "bai12_scenario_comparison.png", "Hình 1. So sánh các kịch bản chính sách tích hợp trong bài 12")

    add_heading(doc, "2. Kiến trúc hệ thống và dữ liệu", 1)
    add_body(
        doc,
        "Cấu trúc dự án được tổ chức theo các lớp rõ ràng: thư mục dashboard chứa giao diện Streamlit; src chứa logic mô hình; data chứa dữ liệu đầu vào; outputs chứa bảng và biểu đồ sinh ra; reports chứa tài liệu báo cáo; scripts chứa công cụ sinh output và đóng gói.",
    )
    add_table(
        doc,
        ["Thư mục/Tệp", "Vai trò"],
        [
            ["dashboard/app.py", "Giao diện web, menu 12 bài, tham số tương tác và hiển thị kết quả"],
            ["src/aideom_vn", "Các module mô hình, solver, trực quan hóa và phân tích"],
            ["data/*.csv", "Dữ liệu vĩ mô, ngành, vùng, dự án và kịch bản"],
            ["outputs/tables", "Bảng kết quả CSV của từng bài"],
            ["outputs/figures", "Biểu đồ PNG dùng cho báo cáo"],
            ["scripts", "Sinh output, kiểm tra chất lượng và đóng gói sản phẩm"],
        ],
    )
    add_callout(
        doc,
        "Nguyên tắc dữ liệu",
        [
            "Các biến khác đơn vị như GDP, xuất khẩu và FDI được chuẩn hóa về năm gốc 2020 = 100 khi cần so sánh xu hướng tương đối.",
            "Các tham số mô phỏng được trình bày như giả định chính sách, không thay thế dự báo chính thức.",
        ],
    )

    add_heading(doc, "3. Mô tả chi tiết 12 chức năng", 1)
    add_table(
        doc,
        ["Bài", "Chức năng chính", "Đầu vào", "Đầu ra"],
        [
            ["1", "Cobb-Douglas mở rộng và TFP", "GDP, vốn, lao động, D, AI, H", "TFP, MAPE, dự báo GDP 2030, phân rã tăng trưởng"],
            ["2", "LP phân bổ ngân sách số", "Ngân sách, ràng buộc tối thiểu, hệ số GDP", "Phân bổ tối ưu và shadow price"],
            ["3", "Ưu tiên 10 ngành", "Trọng số tăng trưởng, việc làm, lan tỏa, xuất khẩu, AI, rủi ro", "Điểm ưu tiên và thứ hạng ngành"],
            ["4", "LP ngành-vùng", "Ngân sách, sàn vùng yếu, trần mỗi vùng", "Phân bổ theo vùng-ngành-hạng mục"],
            ["5", "MIP lựa chọn dự án", "Ngân sách, trần rủi ro, danh mục dự án", "Danh sách dự án được chọn, NPV, chi phí, rủi ro"],
            ["6", "TOPSIS 6 vùng", "Ma trận tiêu chí vùng", "Điểm TOPSIS, trọng số entropy, xếp hạng vùng"],
            ["7", "Pareto đa mục tiêu", "Ngân sách, population, thế hệ, seed", "Pareto frontier, nghiệm thỏa hiệp"],
            ["8", "Tối ưu động 2026-2035", "Ngân sách hằng năm, tỷ lệ phân bổ", "Đường GDP tối ưu và cơ cấu đầu tư"],
            ["9", "Tác động AI tới lao động", "Tỷ lệ ứng dụng AI, ngân sách đào tạo, đầu tư AI", "Việc làm mất, tạo mới, bảo vệ và net job"],
            ["10", "Stochastic programming hai giai đoạn", "Ngân sách và xác suất kịch bản", "Expected value, VSS, EVPI"],
            ["11", "Q-learning chính sách thích nghi", "Episodes, learning rate, discount, epsilon", "Chính sách học được và Q-table"],
            ["12", "Dashboard tích hợp AIDEOM-VN", "Ngân sách kịch bản và trọng số", "So sánh 5 kịch bản, radar KPI, thiết kế hệ thống"],
        ],
    )
    add_picture_if_exists(doc, "bai01_actual_vs_predicted_gdp.png", "Hình 2. Bài 1 - GDP thực tế và giá trị dự báo theo mô hình")
    add_picture_if_exists(doc, "bai02_budget_allocation.png", "Hình 3. Bài 2 - Phân bổ ngân sách số theo nghiệm tối ưu")
    doc.add_page_break()

    add_heading(doc, "4. Tác nhân phân tích kết quả", 1)
    add_body(
        doc,
        "Tác nhân phân tích kết quả là lớp diễn giải nằm ở cuối mỗi bài. Khi người dùng thay đổi tham số, phần này cập nhật theo kết quả hiện tại để chuyển bảng số và biểu đồ thành nhận định chính sách có cấu trúc.",
    )
    add_table(
        doc,
        ["Tầng phân tích", "Nội dung hiển thị"],
        [
            ["Chẩn đoán dữ liệu", "Nêu nguồn dữ liệu, biến quan trọng, đơn vị đo và lưu ý chuẩn hóa"],
            ["Diễn giải chuyên gia", "Giải thích nghiệm tối ưu, thứ hạng, xu hướng, trade-off và ý nghĩa mô hình"],
            ["Hàm ý chính sách", "Rút ra khuyến nghị có thể trình bày khi bảo vệ"],
            ["Kiểm soát rủi ro", "Nhắc các giả định mô phỏng, độ nhạy và giới hạn cần nêu trong báo cáo"],
        ],
    )
    add_body(
        doc,
        "Cách bố trí tác nhân ở cuối mỗi bài giúp hội đồng thấy rõ mối liên hệ giữa thao tác tham số, kết quả mô hình và lập luận ra quyết định. Đây là điểm khác biệt so với việc chỉ nộp notebook hoặc bảng tính tĩnh.",
    )
    add_picture_if_exists(doc, "bai07_pareto_frontier.png", "Hình 4. Bài 7 - Biên Pareto minh họa đánh đổi đa mục tiêu")
    doc.add_page_break()

    add_heading(doc, "5. Hướng dẫn cài đặt và sử dụng", 1)
    add_body(
        doc,
        "Máy cần Python 3.10 trở lên. Dự án chạy local, không bắt buộc Colab. Các bước sử dụng được thiết kế để người chấm có thể mở app nhanh bằng file batch hoặc lệnh Streamlit.",
    )
    add_table(
        doc,
        ["Bước", "Lệnh/Thao tác"],
        [
            ["1", "Giải nén thư mục dự án hoặc file aideom_vn_submission.zip"],
            ["2", "Mở terminal tại thư mục dự án sau khi giải nén"],
            ["3", "Cài thư viện: pip install -r requirements.txt"],
            ["4", "Chạy app: run_app.bat hoặc streamlit run dashboard/app.py"],
            ["5", "Mở trình duyệt tại http://localhost:8501"],
            ["6", "Chọn bài ở sidebar, chỉnh tham số, xem bảng, biểu đồ, tác nhân phân tích và tải CSV"],
        ],
    )
    add_callout(
        doc,
        "Quy trình demo khi bảo vệ",
        [
            "Mở trang chủ để giới thiệu dữ liệu và sơ đồ 12 bài.",
            "Demo bài 1, bài 5, bài 7 và bài 12 để thể hiện đủ dự báo, tối ưu, Pareto và dashboard tích hợp.",
            "Kéo một tham số bất kỳ, chỉ ra bảng, biểu đồ và tác nhân phân tích thay đổi theo kết quả hiện tại.",
        ],
    )

    add_heading(doc, "6. Kết quả đầu ra và kiểm thử", 1)
    add_body(
        doc,
        "Dự án sinh đầy đủ bảng CSV trong outputs/tables và biểu đồ PNG trong outputs/figures. Các output này dùng trực tiếp cho báo cáo Word/PDF và có thể tái tạo bằng script trong thư mục scripts.",
    )
    add_table(
        doc,
        ["Loại output", "Vị trí", "Mục đích"],
        [
            ["Bảng CSV", "outputs/tables", "Lưu nghiệm mô hình, bảng xếp hạng, kịch bản và tham số"],
            ["Biểu đồ PNG", "outputs/figures", "Chèn vào báo cáo và dùng khi thuyết trình"],
            ["Web app", "dashboard/app.py", "Trình diễn mô hình tương tác trước hội đồng"],
            ["Kiểm thử", "tests", "Kiểm tra luồng dữ liệu, mô hình và khả năng sinh output"],
        ],
    )
    add_body(
        doc,
        "Các bước kiểm thử trọng tâm gồm: chạy pytest, sinh lại outputs, mở Streamlit, kiểm tra 12 menu, kiểm tra tham số tương tác, kiểm tra tải CSV và kiểm tra tác nhân phân tích ở cuối từng bài.",
    )

    add_heading(doc, "7. Kết luận", 1)
    add_body(
        doc,
        "AIDEOM-VN đáp ứng định hướng phương án 3: có web app tương tác, 12 menu tương ứng 12 bài, mô hình tính toán trực tiếp, biểu đồ, output lưu file và tác nhân phân tích kết quả. Sản phẩm phù hợp để nộp kèm báo cáo vì vừa có khả năng trình diễn trực quan, vừa giữ được bằng chứng định lượng qua bảng và hình.",
    )
    add_body(
        doc,
        "Khi bảo vệ, trọng tâm nên nhấn mạnh ba điểm: mô hình được tổ chức thành pipeline tái lập; các tham số thể hiện tư duy ra quyết định và phân tích độ nhạy; tác nhân phân tích giúp chuyển kết quả định lượng thành lập luận chính sách rõ ràng.",
    )


def main():
    doc = setup_document()
    add_cover(doc)
    add_toc(doc)
    add_report_content(doc)
    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
